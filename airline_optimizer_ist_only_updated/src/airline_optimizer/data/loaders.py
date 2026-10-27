"""Source readers and explicit source-to-canonical transformations.

CSV is supported by the standard library. XLSX support uses the optional openpyxl
package; workbook rows are never silently treated as individual flight durations.
"""
from __future__ import annotations

import csv
import hashlib
import json
import math
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from airline_optimizer.data.schemas import (
    DataBundle, DataQualityIssue, ODImportance, RankingTransitionWeight,
)
from airline_optimizer.domain.flights import FlightInstance
from airline_optimizer.domain.time_model import to_absolute_minutes


def _norm(value: Any) -> str:
    return str(value).strip().upper() if value is not None else ""


def _as_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip())


def _as_time(value: Any) -> time:
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    if isinstance(value, time):
        return value.replace(second=0, microsecond=0)
    text = str(value).strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time().replace(second=0, microsecond=0)
        except ValueError:
            pass
    raise ValueError(f"unsupported clock time: {value!r}")


def _rows_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames:
            raise ValueError("CSV has no header")
        return list(reader.fieldnames), [dict(row) for row in reader]


def _rows_xlsx(path: Path, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX reading requires openpyxl; install with `pip install openpyxl`") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name is not None:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"worksheet {sheet_name!r} not found; available: {workbook.sheetnames}")
            sheet = workbook[sheet_name]
        else:
            sheet = workbook[workbook.sheetnames[0]]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header_values = next(iterator)
        except StopIteration:
            raise ValueError("worksheet is empty")
        headers = [str(v).strip() if v is not None else "" for v in header_values]
        rows: list[dict[str, Any]] = []
        for values in iterator:
            if all(v is None for v in values):
                continue
            rows.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
        return headers, rows
    finally:
        workbook.close()


def read_table(path: str | Path, sheet_name: str | None = None) -> tuple[list[str], list[dict[str, Any]]]:
    """Read a CSV/XLSX table without assigning business meaning to its fields."""
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(source)
    suffix = source.suffix.lower()
    if suffix == ".csv":
        headers, rows = _rows_csv(source)
    elif suffix in {".xlsx", ".xlsm"}:
        headers, rows = _rows_xlsx(source, sheet_name)
    else:
        raise ValueError(f"unsupported source format: {suffix}")
    return headers, rows


def _record_source(bundle: DataBundle, path: Path, headers: list[str], rows: list[dict[str, Any]], sheet: str | None) -> None:
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    bundle.source_manifest[str(path.resolve())] = {
        "sha256": digest, "size_bytes": path.stat().st_size, "row_count": len(rows),
        "columns": headers, "sheet": sheet,
    }


def _issue(bundle: DataBundle, code: str, severity: str, path: Path, message: str, row: int | None = None, field: str | None = None) -> None:
    bundle.issues.append(DataQualityIssue(code, severity, path.name, message, row, field))


def load_baseline_flights(
    path: str | Path, bundle: DataBundle | None = None, *,
    duration_column: str | None = None, allow_duration_inference: bool = True,
) -> DataBundle:
    """Load baseline CSV/XLSX into canonical flights.

    If no explicit duration column is supplied, duration is inferred from baseline
    departure/arrival clock times with a single midnight rollover. Such rows receive
    a warning because this is not equivalent to independently validated block time.
    """
    result = bundle or DataBundle()
    source = Path(path)
    headers, rows = read_table(source)
    _record_source(result, source, headers, rows, None)
    required = {"flight_no", "origin", "destination", "departure_date", "departure_time", "arrival_time"}
    missing = required - set(headers)
    if missing:
        _issue(result, "missing_required_columns", "error", source, f"missing columns: {sorted(missing)}")
        return result
    seen: dict[str, tuple[Any, ...]] = {}
    for index, row in enumerate(rows, start=2):
        try:
            number = str(row["flight_no"]).strip()
            origin, destination = _norm(row["origin"]), _norm(row["destination"])
            flight_day = _as_date(row["departure_date"])
            dep_time, arr_time = _as_time(row["departure_time"]), _as_time(row["arrival_time"])
            dep_abs = to_absolute_minutes(flight_day, dep_time)
            if duration_column and duration_column in row and row[duration_column] not in (None, ""):
                duration = int(row[duration_column])
                duration_source = "explicit_column"
            else:
                if not allow_duration_inference:
                    raise ValueError("no validated duration column; inference is disabled")
                dep_m = dep_time.hour * 60 + dep_time.minute
                arr_m = arr_time.hour * 60 + arr_time.minute
                duration = (arr_m - dep_m) % 1440
                if duration == 0:
                    raise ValueError("departure and arrival clocks imply zero duration; explicit duration required")
                duration_source = "inferred_from_baseline_clock_times"
            if duration <= 0 or duration > 24 * 60:
                raise ValueError(f"flight duration outside accepted range: {duration}")
            if not number or not origin or not destination or origin == destination:
                raise ValueError("missing flight number/airport or identical origin and destination")
            flight_id = f"{flight_day.isoformat()}:{number}:{origin}-{destination}"
            fingerprint = (number, flight_day.isoformat(), origin, destination, dep_abs, duration)
            if flight_id in seen:
                if seen[flight_id] == fingerprint:
                    _issue(result, "duplicate_flight_row", "warning", source, f"exact duplicate for {flight_id}", index)
                    continue
                _issue(result, "conflicting_flight_identity", "error", source, f"conflicting rows map to flight_id {flight_id}", index)
                continue
            seen[flight_id] = fingerprint
            result.flights.append(FlightInstance(flight_id, number, flight_day, origin, destination, dep_abs, duration, (f"{source.name}:{index}",)))
            if duration_source != "explicit_column":
                _issue(result, "duration_inferred", "warning", source, "duration inferred from baseline clock times; independently validate before official scoring", index, "arrival_time")
        except (TypeError, ValueError, OverflowError) as exc:
            _issue(result, "invalid_flight_row", "error", source, str(exc), index)
    return result


def load_od_importance(
    path: str | Path, bundle: DataBundle | None = None, *, sheet_name: str = "Yolcu Verisi Son 1 Yıl",
) -> DataBundle:
    result = bundle or DataBundle()
    source = Path(path)
    headers, rows = read_table(source, sheet_name if source.suffix.lower() in {".xlsx", ".xlsm"} else None)
    _record_source(result, source, headers, rows, sheet_name if source.suffix.lower() in {".xlsx", ".xlsm"} else None)
    aliases = {"origin": "Orig Airport Code", "destination": "Dest Airport Code", "importance": "OD importance factor"}
    missing = set(aliases.values()) - set(headers)
    if missing:
        _issue(result, "missing_required_columns", "error", source, f"missing columns: {sorted(missing)}")
        return result
    staged: dict[tuple[str, str], list[tuple[float, str, int]]] = {}
    for index, row in enumerate(rows, start=2):
        try:
            origin, destination = _norm(row[aliases["origin"]]), _norm(row[aliases["destination"]])
            importance = float(row[aliases["importance"]])
            if not origin or not destination or origin == destination or not math.isfinite(importance) or importance < 0:
                raise ValueError("invalid OD codes or importance value")
            staged.setdefault((origin, destination), []).append((importance, f"{source.name}:{index}", index))
        except (TypeError, ValueError) as exc:
            _issue(result, "invalid_od_importance_row", "error", source, str(exc), index)
    for key, values in staged.items():
        distinct = {value for value, _, _ in values}
        if len(distinct) > 1:
            _issue(result, "conflicting_od_importance", "error", source, f"OD {key} has conflicting importance values: {sorted(distinct)}")
            continue
        if len(values) > 1:
            _issue(result, "duplicate_od_importance", "warning", source, f"OD {key} appears {len(values)} times with the same value")
        value = values[0][0]
        result.od_importance[key] = ODImportance(key[0], key[1], value, tuple(item[1] for item in values))
    return result


def load_ranking_weights(
    path: str | Path, bundle: DataBundle | None = None, *, sheet_name: str = "Sheet1",
) -> DataBundle:
    result = bundle or DataBundle()
    source = Path(path)
    headers, rows = read_table(source, sheet_name if source.suffix.lower() in {".xlsx", ".xlsm"} else None)
    _record_source(result, source, headers, rows, sheet_name if source.suffix.lower() in {".xlsx", ".xlsm"} else None)
    required = {"Rakip Sayısı", "İlk Rank", "Son Rank", "Weight"}
    missing = required - set(headers)
    if missing:
        _issue(result, "missing_required_columns", "error", source, f"missing columns: {sorted(missing)}")
        return result
    seen: dict[tuple[int, int, int], float] = {}
    for index, row in enumerate(rows, start=2):
        try:
            count, first, last = int(row["Rakip Sayısı"]), int(row["İlk Rank"]), int(row["Son Rank"])
            weight = float(row["Weight"])
            if count < 1 or first < 1 or last < 1 or first > count or last > count or not math.isfinite(weight):
                raise ValueError("invalid competitor count, rank, or weight")
            key = (count, first, last)
            if key in seen and seen[key] != weight:
                _issue(result, "conflicting_ranking_weight", "error", source, f"key {key} has conflicting weights", index)
                continue
            if key in seen:
                _issue(result, "duplicate_ranking_weight", "warning", source, f"duplicate key {key}", index)
                continue
            seen[key] = weight
            result.ranking_weights.append(RankingTransitionWeight(count, first, last, weight, f"{source.name}:{index}"))
        except (TypeError, ValueError) as exc:
            _issue(result, "invalid_ranking_weight_row", "error", source, str(exc), index)
    return result


def load_flight_pairs_metadata(path: str | Path, bundle: DataBundle | None = None, *, sheet_name: str = "W26 Flight Pair 20260227") -> DataBundle:
    """Preserve flight-pair source rows as metadata; do not use Pair as a hard pairing."""
    result = bundle or DataBundle()
    source = Path(path)
    headers, rows = read_table(source, sheet_name if source.suffix.lower() in {".xlsx", ".xlsm"} else None)
    _record_source(result, source, headers, rows, sheet_name if source.suffix.lower() in {".xlsx", ".xlsm"} else None)
    required = {"FlNo", "Orig", "Dest", "Pair"}
    missing = required - set(headers)
    if missing:
        _issue(result, "missing_required_columns", "error", source, f"missing columns: {sorted(missing)}")
        return result
    result.raw_tables["flight_pairs_metadata"] = rows
    # Pair is intentionally retained only as source metadata, not as a connection restriction.
    return result


def write_quality_report(bundle: DataBundle, path: str | Path) -> Path:
    """Write a deterministic JSON report with issues and source fingerprints."""
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "status": "ERRORS" if bundle.has_errors else ("WARNINGS" if bundle.issues else "OK"),
        "flight_count": len(bundle.flights), "od_importance_count": len(bundle.od_importance),
        "ranking_weight_count": len(bundle.ranking_weights),
        "source_manifest": bundle.source_manifest,
        "issues": [
            {"code": i.code, "severity": i.severity, "source_file": i.source_file, "message": i.message, "row_number": i.row_number, "field": i.field}
            for i in bundle.issues
        ],
    }
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True), encoding="utf-8")
    return target
